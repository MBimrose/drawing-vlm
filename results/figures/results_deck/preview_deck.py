"""Approximate slide previews from shape geometry (no LibreOffice on the cluster): pictures pasted at their boxes;
rectangles drawn with their fill and outline (theme colours resolved from the deck's theme); lines drawn as lines;
text wrapped in Nimbus Sans with its alignment, size, weight and colour. A red box marks text whose estimated height
overflows its box.

    PYTHONPATH=<dir with python-pptx> python preview_deck.py deck.pptx out_prefix   -> out_prefix_<n>.png
"""
import io
import re
import sys
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
BOLD = "/usr/share/fonts/urw-base35/NimbusSans-Bold.otf"; REG = "/usr/share/fonts/urw-base35/NimbusSans-Regular.otf"
src, prefix = sys.argv[1], sys.argv[2]
pr = Presentation(src); DPI = 110
W, H = int(pr.slide_width / 914400 * DPI), int(pr.slide_height / 914400 * DPI)
px = lambda e: int(e / 914400 * DPI)

z = zipfile.ZipFile(src); theme = z.read([n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n)][0]).decode()
TH = {k: v for k, v in re.findall(r"<a:(dk1|lt1|dk2|lt2|accent\d|hlink|folHlink)><a:(?:srgbClr val|sysClr [^>]*lastClr)=\"([0-9A-F]{6})\"", theme)}
ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


def colour(el, default=None):
    """First solidFill colour under an XML element (srgb or scheme), as an RGB tuple."""
    if el is None:
        return default
    f = el.find(A + "solidFill")
    if f is None:
        return default
    c = f.find(A + "srgbClr")
    if c is not None:
        v = c.get("val")
    else:
        c = f.find(A + "schemeClr")
        if c is None:
            return default
        v = TH.get(ALIAS.get(c.get("val"), c.get("val")))
    return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4)) if v else default


for si, s in enumerate(pr.slides):
    im = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(im)
    for sh in s.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        if sh.shape_type == 13:
            pic = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA").resize((max(w, 1), max(h, 1))); im.paste(pic, (x, y), pic)
            continue
        spPr = sh.element.find(".//" + "{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
        geom = spPr.find(A + "prstGeom").get("prst") if spPr is not None and spPr.find(A + "prstGeom") is not None else ""
        ln = spPr.find(A + "ln") if spPr is not None else None
        lc = colour(ln, None); lw = max(1, round(int(ln.get("w", 12700)) / 12700 * DPI / 72)) if ln is not None and ln.get("w") else 1
        if geom == "line":
            xf = sh.element.xpath(".//a:xfrm")[0]; fh, fv = xf.get("flipH") == "1", xf.get("flipV") == "1"
            d.line((x + (w if fh else 0), y + (h if fv else 0), x + (0 if fh else w), y + (0 if fv else h)), fill=lc or (0, 0, 0), width=lw)
            continue
        fc = colour(spPr, None)
        if fc or lc:
            box = (x, y, x + w, y + h)
            if geom == "roundRect":
                d.rounded_rectangle(box, radius=int(0.08 * DPI), fill=fc, outline=lc, width=lw)
            else:
                d.rectangle(box, fill=fc, outline=lc, width=lw)
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        lines = []   # (text, font, colour, align)
        for p in sh.text_frame.paragraphs:
            if not p.runs:
                continue
            al = p.alignment
            bullet = p._pPr is not None and (p._pPr.find(A + "buChar") is not None or p._pPr.find(A + "buAutoNum") is not None)
            segs = []
            for r in p.runs:
                sz = r.font.size.pt if r.font.size else 32
                f = ImageFont.truetype(BOLD if (r.font.bold or sh.is_placeholder) else REG, int(sz * DPI / 72))
                segs.append((r.text, f, colour(r._r.find(A + "rPr"), (0, 0, 0)), sz))
                if r._r.getnext() is not None and r._r.getnext().tag == A + "br":
                    segs.append(("\n", f, None, sz))
            cur, curw = [], 0
            words = [("• " if bullet else "", segs[0][1], segs[0][2], segs[0][3])] if bullet else []
            for t, f, c, sz in segs:
                if t == "\n":
                    words.append(("\n", f, c, sz)); continue
                for wd in re.split(r"(?<= )", t):
                    words.append((wd, f, c, sz))
            line, lw_ = [], 0
            for wd, f, c, sz in words:
                if wd == "\n":
                    lines.append((line, al)); line, lw_ = [], 0; continue
                ww = d.textlength(wd, font=f)
                if lw_ + ww > w and line:
                    lines.append((line, al)); line, lw_ = [], 0
                line.append((wd, f, c, sz)); lw_ += ww
            if line:
                lines.append((line, al))
        hts = [int(max(sz for _, _, _, sz in ln_) * DPI / 72 * 1.18) for ln_, _ in lines]
        total = sum(hts); anchor = sh.text_frame._txBody.find(A + "bodyPr").get("anchor") if sh.text_frame._txBody.find(A + "bodyPr") is not None else None
        cy = y + (h - total) // 2 if anchor == "ctr" or sh.is_placeholder and False else (y + h - total if anchor == "b" else y)
        if anchor == "ctr":
            cy = y + (h - total) // 2
        for (ln_, al), lh in zip(lines, hts):
            tw = sum(d.textlength(t, font=f) for t, f, _, _ in ln_)
            cx = x + (w - tw) / 2 if al == PP_ALIGN.CENTER else (x + w - tw if al == PP_ALIGN.RIGHT else x)
            for t, f, c, _ in ln_:
                d.text((cx, cy), t, fill=c or (0, 0, 0), font=f); cx += d.textlength(t, font=f)
            cy += lh
        if total > h + 4:
            d.rectangle((x, y, x + w, y + h), outline="red", width=3)
    im.save(f"{prefix}_{si + 1}.png")
print(len(pr.slides), "slides")
